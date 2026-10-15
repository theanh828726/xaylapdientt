"""Tạo 6 mẫu biểu .docx mặc định vào app/templates/ (chạy: python scripts/tao_mau_bieu.py).

Mẫu dùng thẻ docxtpl ({{ biến }}, {%tr for ... %}). Có thể mở bằng Word để chỉnh bố cục/chữ rồi tải lên lại
trong màn hình Quản trị → Biểu mẫu; giữ nguyên các thẻ {{ }} và {% %}.
Đây là mẫu mặc định do hệ thống tạo, cần đối chiếu biểu mẫu hiện hành (TT 10/2021/TT-BXD) trước khi phát hành chính thức.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent.parent / "app" / "templates"
LUU_Y = ("Mẫu mặc định do hệ thống TuyếnSố tạo — cần đối chiếu với biểu mẫu hiện hành "
         "(TT 10/2021/TT-BXD) trước khi phát hành chính thức.")


def moi() -> Document:
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(4)
    return d


def dong(d, text="", *, bold=False, italic=False, center=False, size=None, underline=False, after=None):
    p = d.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold, r.italic, r.underline = bold, italic, underline
    if size:
        r.font.size = Pt(size)
    return p


def nhan_gia_tri(d, nhan, gia_tri):
    """'Nhãn: ' in đậm + giá trị (mỗi thẻ {{ }} nằm trọn trong một run để docxtpl nhận diện)."""
    p = d.add_paragraph()
    p.add_run(nhan).bold = True
    p.add_run(gia_tri)
    return p


def quoc_hieu(d):
    dong(d, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, center=True, after=0)
    dong(d, "Độc lập – Tự do – Hạnh phúc", bold=True, center=True, underline=True, after=10)


def bang(d, tieu_de, widths, hang_lap, ten_bien):
    """Bảng có dòng lặp docxtpl: dòng {%tr for %} → dòng dữ liệu → dòng {%tr endfor %}."""
    t = d.add_table(rows=1, cols=len(tieu_de))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(tieu_de):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(h).bold = True
    bat_dau = t.add_row().cells
    bat_dau[0].text = f"{{%tr for {ten_bien} in {hang_lap} %}}"
    du_lieu = t.add_row().cells
    for i, v in enumerate(widths):
        du_lieu[i].text = v
    ket_thuc = t.add_row().cells
    ket_thuc[0].text = "{%tr endfor %}"
    return t


def chu_ky(d, cot):
    t = d.add_table(rows=2, cols=len(cot))
    for i, ten in enumerate(cot):
        a, b = t.rows[0].cells[i], t.rows[1].cells[i]
        a.text, b.text = "", "(Ký, ghi rõ họ tên)"
        a.paragraphs[0].add_run(ten).bold = True
        for c in (a, b):
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        b.paragraphs[0].runs[0].italic = True
        for _ in range(4):
            b.add_paragraph("")
    return t


def chan_trang(d):
    dong(d, "")
    dong(d, LUU_Y, italic=True, size=9)


def nghiem_thu() -> Document:
    d = moi()
    quoc_hieu(d)
    dong(d, "BIÊN BẢN NGHIỆM THU CÔNG VIỆC XÂY DỰNG", bold=True, center=True, size=15, after=0)
    dong(d, "Số: {{ so }}", center=True, after=10)
    nhan_gia_tri(d, "Công trình: ", "{{ du_an.ten }}")
    nhan_gia_tri(d, "Chủ đầu tư: ", "{{ du_an.chu_dau_tu }}")
    nhan_gia_tri(d, "Công tác nghiệm thu: ", "{{ hang_muc }}")
    nhan_gia_tri(d, "Vị trí: ", "{{ vi_tri_text }}")
    nhan_gia_tri(d, "Ngày hoàn thành: ", "{{ ngay_hoan_thanh }}")
    nhan_gia_tri(d, "Thời gian nghiệm thu: ", "{{ ngay.gio_phut }}, ngày {{ ngay.ngay }} tháng {{ ngay.thang }} năm {{ ngay.nam }}")
    dong(d, "I. Căn cứ nghiệm thu", bold=True)
    dong(d, "- Hồ sơ thiết kế, chỉ dẫn kỹ thuật và hợp đồng của dự án.")
    dong(d, "- Hồ sơ nghiệm thu kèm theo: {{ phu_luc }}.")
    dong(d, "II. Các vị trí được nghiệm thu", bold=True)
    bang(d, ["STT", "Vị trí", "Km", "Loại móng", "Khối lượng hố móng (m³)"],
         ["{{ v.stt }}", "{{ v.ten }}", "{{ v.km }}", "{{ v.loai_mong }}", "{{ v.khoi_luong }}"], "danh_sach", "v")
    dong(d, "")
    dong(d, "III. Thành phần trực tiếp nghiệm thu", bold=True)
    bang(d, ["Họ và tên", "Chức danh", "Đơn vị"], ["{{ t.ho_ten }}", "{{ t.chuc_danh }}", "{{ t.don_vi }}"], "thanh_phan", "t")
    dong(d, "")
    dong(d, "IV. Kết luận", bold=True)
    dong(d, "Chấp nhận nghiệm thu {{ hang_muc }} tại {{ vi_tri_text }} để chuyển sang công việc tiếp theo.")
    dong(d, "")
    chu_ky(d, ["ĐẠI DIỆN NHÀ THẦU THI CÔNG", "ĐẠI DIỆN TƯ VẤN GIÁM SÁT", "ĐẠI DIỆN CHỦ ĐẦU TƯ"])
    chan_trang(d)
    return d


def giay_moi() -> Document:
    d = moi()
    quoc_hieu(d)
    dong(d, "GIẤY MỜI NGHIỆM THU", bold=True, center=True, size=15, after=0)
    dong(d, "Số: {{ so }}", center=True, after=10)
    dong(d, "Kính gửi: Tư vấn giám sát và {{ du_an.chu_dau_tu }}")
    dong(d, "Nhà thầu thi công trân trọng kính mời Quý đơn vị tham dự nghiệm thu công việc xây dựng thuộc công trình {{ du_an.ten }}:")
    nhan_gia_tri(d, "Công tác: ", "{{ cong_viec }}")
    nhan_gia_tri(d, "Các vị trí: ", "{{ vi_tri_text }}")
    nhan_gia_tri(d, "Thời gian dự kiến: ngày ", "{{ ngay_nghiem_thu }}")
    dong(d, "Địa điểm: tại hiện trường các vị trí nêu trên.")
    dong(d, "")
    dong(d, "Rất mong Quý đơn vị sắp xếp tham dự đúng giờ.")
    dong(d, "")
    dong(d, "Ngày {{ ngay.ngay }} tháng {{ ngay.thang }} năm {{ ngay.nam }}", italic=True, center=True)
    dong(d, "{{ nguoi_lap.chuc_danh }} – {{ nguoi_lap.ho_ten }}", bold=True, center=True)
    chan_trang(d)
    return d


def chiet_tinh() -> Document:
    d = moi()
    quoc_hieu(d)
    dong(d, "BẢNG CHIẾT TÍNH KHỐI LƯỢNG HỐ MÓNG", bold=True, center=True, size=15, after=0)
    dong(d, "Số: {{ so }}", center=True, after=10)
    nhan_gia_tri(d, "Công trình: ", "{{ du_an.ten }}")
    nhan_gia_tri(d, "Kèm theo công tác: ", "{{ cong_viec }}")
    dong(d, "Công thức: A1 = A + 0,4; A2 = A1 + 2Kh; S1 = A1·B1; S2 = A2·B2; V = h/3·(S1 + S2 + √(S1·S2)).", italic=True, size=11)
    bang(d, ["STT", "Vị trí", "Loại móng", "A", "B", "h", "K", "S1 (đáy) m²", "S2 (miệng) m²", "V (m³)"],
         ["{{ r.stt }}", "{{ r.vi_tri }}", "{{ r.loai_mong }}", "{{ r.a }}", "{{ r.b }}", "{{ r.h }}", "{{ r.k }}", "{{ r.s_day }}", "{{ r.s_mieng }}", "{{ r.v }}"],
         "rows", "r")
    dong(d, "")
    nhan_gia_tri(d, "Tổng khối lượng: ", "{{ tong_v }} m³")
    dong(d, "")
    chu_ky(d, ["NGƯỜI LẬP", "CHỈ HUY TRƯỞNG"])
    chan_trang(d)
    return d


def de_trinh() -> Document:
    d = moi()
    quoc_hieu(d)
    dong(d, "ĐỀ TRÌNH CHẤP THUẬN NHÀ CUNG CẤP VẬT LIỆU", bold=True, center=True, size=15, after=10)
    dong(d, "Kính gửi: {{ du_an.chu_dau_tu }}")
    dong(d, "Công trình: {{ du_an.ten }}")
    dong(d, "Nhà thầu đề nghị Chủ đầu tư xem xét chấp thuận nhà cung cấp vật liệu sau:")
    nhan_gia_tri(d, "Vật liệu: ", "{{ vat_lieu.ten }} (lô {{ vat_lieu.ma }})")
    nhan_gia_tri(d, "Nhà cung cấp: ", "{{ vat_lieu.nha_cung_cap }}")
    nhan_gia_tri(d, "Ngày giao: ", "{{ vat_lieu.ngay_giao }}")
    nhan_gia_tri(d, "Đã sử dụng tại: ", "{{ vat_lieu.dung_tai }}")
    nhan_gia_tri(d, "Đề nghị có văn bản chấp thuận trước ngày: ", "{{ vat_lieu.han }}")
    dong(d, "Hồ sơ năng lực, chứng chỉ CO/CQ và kết quả thí nghiệm của lô vật liệu được gửi kèm theo đề trình này.")
    dong(d, "")
    dong(d, "Ngày {{ ngay.ngay }} tháng {{ ngay.thang }} năm {{ ngay.nam }}", italic=True, center=True)
    dong(d, "{{ nguoi_lap.chuc_danh }} – {{ nguoi_lap.ho_ten }}", bold=True, center=True)
    chan_trang(d)
    return d


def de_xuat() -> Document:
    d = moi()
    quoc_hieu(d)
    dong(d, "PHIẾU ĐỀ XUẤT VẬT TƯ", bold=True, center=True, size=15, after=0)
    dong(d, "Số: {{ ma }}", center=True, after=10)
    nhan_gia_tri(d, "Dự án: ", "{{ du_an.ten }}")
    nhan_gia_tri(d, "Công trình/đoạn tuyến: ", "{{ cong_trinh }}")
    nhan_gia_tri(d, "Người lập: ", "{{ nguoi_lap_ten }}")
    nhan_gia_tri(d, "Ngày lập: ", "{{ ngay_lap }}")
    nhan_gia_tri(d, "Trạng thái: ", "{{ trang_thai }}")
    nhan_gia_tri(d, "Giá hợp đồng: ", "{{ gia_hop_dong }}")
    nhan_gia_tri(d, "Giá khoán: ", "{{ gia_khoan }}")
    dong(d, "Danh mục vật tư đề xuất", bold=True)
    bang(d, ["STT", "Mã", "Vật tư", "ĐVT", "Định mức", "Đề xuất", "Tồn kho"],
         ["{{ r.stt }}", "{{ r.ma_ck }}", "{{ r.ten }}", "{{ r.dvt }}", "{{ r.theo_dinh_muc }}", "{{ r.de_xuat }}", "{{ r.ton_kho }}"],
         "dong", "r")
    dong(d, "")
    dong(d, "Chuỗi phê duyệt", bold=True)
    bang(d, ["Bước", "Vai trò", "Người duyệt", "Trạng thái", "Thời gian", "Ghi chú"],
         ["{{ b.stt }}", "{{ b.vai_tro }}", "{{ b.nguoi }}", "{{ b.trang_thai }}", "{{ b.thoi_gian }}", "{{ b.ly_do }}"],
         "chuoi_duyet", "b")
    dong(d, "")
    chu_ky(d, ["NGƯỜI LẬP", "CHỈ HUY TRƯỞNG", "PHÊ DUYỆT"])
    chan_trang(d)
    return d


def dieu_chinh() -> Document:
    d = moi()
    quoc_hieu(d)
    dong(d, "PHIẾU ĐIỀU CHỈNH THIẾT KẾ", bold=True, center=True, size=15, after=0)
    dong(d, "Số: {{ ma }}", center=True, after=10)
    nhan_gia_tri(d, "Dự án: ", "{{ du_an.ten }}")
    nhan_gia_tri(d, "Ngày lập: ", "{{ ngay_lap }}")
    nhan_gia_tri(d, "Trạng thái: ", "{{ trang_thai }}")
    nhan_gia_tri(d, "Nội dung/lý do điều chỉnh: ", "{{ ly_do }}")
    dong(d, "")
    dong(d, "Chuỗi phê duyệt", bold=True)
    bang(d, ["Bước", "Vai trò", "Người duyệt", "Trạng thái", "Thời gian", "Ghi chú"],
         ["{{ b.stt }}", "{{ b.vai_tro }}", "{{ b.nguoi }}", "{{ b.trang_thai }}", "{{ b.thoi_gian }}", "{{ b.ly_do }}"],
         "chuoi_duyet", "b")
    dong(d, "")
    chu_ky(d, ["NGƯỜI LẬP", "TƯ VẤN THIẾT KẾ", "PHÊ DUYỆT"])
    chan_trang(d)
    return d


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for ten, ham in {"nghiem_thu_cong_viec": nghiem_thu, "de_xuat_vat_tu": de_xuat, "phieu_dieu_chinh": dieu_chinh,
                     "giay_moi_nghiem_thu": giay_moi, "bang_chiet_tinh": chiet_tinh, "de_trinh_nha_cung_cap": de_trinh}.items():
        ham().save(OUT / f"{ten}.docx")
        print("đã tạo", OUT / f"{ten}.docx")
