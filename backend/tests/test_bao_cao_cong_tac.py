"""Báo cáo hiện trường theo công tác (mobile) và biên bản của một công tác. Dùng dự án 2 để không đổi số liệu dự án 1."""
import io

import pytest_asyncio
from docx import Document
from sqlalchemy import select

from .conftest import new_id


def doc_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


async def cong_tac_id(vi_tri: str, ma: str, du_an: str = "DEMO-BL110") -> str:
    from app.db import SessionLocal
    from app.models import CongTac, CongTacMa, DuAn, ViTriCot
    async with SessionLocal() as s:
        da = await s.scalar(select(DuAn.id).where(DuAn.ma == du_an))
        return str(await s.scalar(select(CongTac.id).join(ViTriCot, ViTriCot.id == CongTac.vi_tri_id).where(
            CongTac.du_an_id == da, ViTriCot.ten == vi_tri, CongTac.ma == CongTacMa(ma))))


@pytest_asyncio.fixture(scope="module", loop_scope="session")
async def p2(client, H):
    du_an = (await client.get("/api/du-an", headers=H["giamdoc"])).json()
    h = {**H["giamdoc"], "X-Du-An": next(d["id"] for d in du_an if d["ma"] == "DEMO-BL110")}
    cot = {c["ten"]: c for c in (await client.get("/api/vi-tri-cot", headers=h)).json()}
    return h, cot


async def test_form_bao_cao_vi_tri_48_theo_thiet_ke(client, H):
    cot = {c["ten"]: c for c in (await client.get("/api/vi-tri-cot", headers=H["tam"])).json()}
    f = (await client.get(f"/api/vi-tri-cot/{cot['VT48']['id']}/bao-cao-form", headers=H["tam"])).json()
    assert f["viTri"]["ten"] == "VT48" and f["viTri"]["kmText"] == "Km 2+550" and f["viTri"]["loaiMong"] == "MT-7a-14-2,8"
    assert f["khoiLuongThietKe"] == {"daoMong": 62.0, "dvt": "m³", "datCap": "Cấp I"}        # đúng kiểm chứng 62,0 m³ trong CLAUDE.md
    assert f["congTacMacDinh"] == "dao_mong"
    khoa = {c["ma"]: c["khoa"] for c in f["congTac"]}
    assert khoa["dao_mong"] is False and khoa["bt_lot"] is True and khoa["dung_cot"] is True
    assert next(c for c in f["congTac"] if c["ma"] == "bt_lot")["lyDoKhoa"] == "Mở khi Đào móng đã nghiệm thu"
    assert f["canhBao"] == ["Chưa có BB bàn giao tim mốc VT 48. Chụp ảnh mốc trước khi đào để hồ sơ không thiếu."]
    assert "Tổ 3 · Lê Văn Tâm" in f["toDoi"] and f["anhToiThieu"] == 3 and "Máy đào 0,8 m³" in f["mayThiCong"]
    assert 105 < f["viTri"]["lon"] < 106


async def test_bao_cao_cong_tac_khoa_theo_tien_quyet_va_cap_nhat_tien_do(client, H, p2):
    h, cot = p2
    base = {"viTri": "VT1", "toDoi": "Tổ 3 · Lê Văn Tâm", "khoiLuong": 61.5, "nhanLuc": 5, "thoiTiet": "Nắng, 31°C", "mayThiCong": ["Máy đào 0,8 m³"]}
    r = await client.post("/api/bao-cao-hien-truong", headers=h, json={**base, "congTacMa": "bt_lot"})
    assert r.status_code == 409 and "R02" in r.json()["detail"]              # bê tông lót chưa mở vì đào móng chưa nghiệm thu
    r = await client.post("/api/bao-cao-hien-truong", headers=h, json={**base, "congTacMa": "dao_mong", "hoanThanh": True, "clientId": new_id()})
    assert r.status_code == 201, r.text
    b = r.json()
    assert b["congTacMa"] == "dao_mong" and b["khoiLuong"] == 61.5 and b["noiDung"].endswith("hoàn thành, chờ nghiệm thu") and b["mayThiCong"] == ["Máy đào 0,8 m³"]
    d = (await client.get(f"/api/vi-tri-cot/{cot['VT1']['id']}", headers=h)).json()
    assert d["lichSu"][0]["trangThai"] == "dang" and d["lichSu"][0]["nhan"] == "Chờ NT"          # đã làm xong, chờ nghiệm thu
    r = await client.post("/api/bao-cao-hien-truong", headers=h, json={**base, "congTacMa": "dao_mong"})
    assert r.status_code == 409 and "đã hoàn thành" in r.json()["detail"]
    assert (await client.post("/api/bao-cao-hien-truong", headers=h, json={**base, "congTacMa": "bt_lot"})).status_code == 409


async def test_bao_cao_khong_hoan_thanh_chi_dang_lam(client, H, p2):
    h, cot = p2
    r = await client.post("/api/bao-cao-hien-truong", headers=h, json={"viTri": "VT2", "congTacMa": "dao_mong", "khoiLuong": 20, "toDoi": "Tổ 3"})
    assert r.status_code == 201 and "đang thực hiện" in r.json()["noiDung"]
    d = (await client.get(f"/api/vi-tri-cot/{cot['VT2']['id']}", headers=h)).json()
    assert d["lichSu"][0]["nhan"] == "Đang làm" and d["chips"][0]["text"] == "Đang thi công móng"


async def test_bien_ban_mot_cong_tac_du_dieu_kien_va_noi_dung(client, H, p2):
    h, _ = p2
    ct_id, chua = await cong_tac_id("VT1", "dao_mong"), await cong_tac_id("VT3", "dao_mong")
    assert (await client.post(f"/api/cong-tac/{chua}/bien-ban", headers=h)).status_code == 409            # chưa hoàn thành
    assert (await client.post(f"/api/cong-tac/{ct_id}/bien-ban", headers={**h, "Authorization": H["ketoan"]["Authorization"]})).status_code == 403
    r = await client.post(f"/api/cong-tac/{ct_id}/bien-ban", headers=h)
    assert r.status_code == 201, r.text
    bb = r.json()
    assert bb["so"].startswith("BB-NT-") and bb["mauNguon"] == "mac-dinh" and bb["congTacIds"] == [ct_id]
    text = doc_text((await client.get(f"/api/bien-ban/{bb['id']}/tai-ve", headers=h)).content)
    for can in (bb["so"], "Đào móng", "VT 1", "ĐZ 110kV Nghi Sơn", "PL1 · Phần móng", "Km 0+000", "MT-8a-20-3,7", "142,7", "Đại diện tư vấn giám sát"):
        assert can in text, can
    assert "{{" not in text and "{%" not in text
    again = await client.post(f"/api/cong-tac/{ct_id}/bien-ban", headers=h)
    assert again.status_code == 200 and again.json()["id"] == bb["id"]               # không phát hành trùng
    moi = await client.post(f"/api/cong-tac/{ct_id}/bien-ban?tao_lai=true", headers=h)
    assert moi.status_code == 201 and moi.json()["so"] != bb["so"]
    ds = (await client.get(f"/api/cong-tac/{ct_id}/bien-ban", headers=h)).json()
    assert [b["so"] for b in ds] == [bb["so"], moi.json()["so"]]


async def test_bien_ban_cong_tac_bi_chan_khong_duoc_sinh(client, H):
    vt46 = await cong_tac_id("VT46", "bt_lot", "DEMO-LO471")
    r = await client.post(f"/api/cong-tac/{vt46}/bien-ban", headers=H["kythuat"])
    assert r.status_code == 409 and "R04" in r.json()["detail"] and "R03" in r.json()["detail"]
    vt38 = await cong_tac_id("VT38", "dung_cot", "DEMO-LO471")
    assert "R06" in (await client.post(f"/api/cong-tac/{vt38}/bien-ban", headers=H["kythuat"])).json()["detail"]
